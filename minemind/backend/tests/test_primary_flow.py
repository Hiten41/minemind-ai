import io
from unittest.mock import AsyncMock, patch

import fitz
from docx import Document
from fastapi.testclient import TestClient

from api import query
from main import app
from services import document_lifecycle as lifecycle
from test_query_contract import AnswerClient


def test_primary_flow_isolated_providers(store):
    pdf = fitz.open()
    pdf.new_page().insert_text((50, 50), 'Pump pressure must not exceed 42 kPa.')
    digital = pdf.tobytes()
    pdf.new_page()
    mixed = pdf.tobytes()
    pdf.close()
    docx = Document()
    docx.add_table(rows=1, cols=1).cell(0, 0).text = 'Pump pressure limit is 42 kPa.'
    stream = io.BytesIO()
    docx.save(stream)
    with TestClient(app) as client, patch.object(lifecycle, 'ingest_document', AsyncMock(return_value=1)), patch.object(lifecycle, 'delete_dataset', AsyncMock(return_value=True)), patch('services.retrieval.query_memory', AsyncMock(return_value=[])):
        registered = client.post('/api/auth/register', json={'name': 'Flow Test', 'email': 'flow@example.com', 'mobile': '+917777777777', 'password': 'a-long-test-password'})
        assert registered.status_code == 200
        login = client.post('/api/auth/login', json={'identifier': 'flow@example.com', 'password': 'a-long-test-password'})
        headers = {'Authorization': 'Bearer ' + login.json()['token']}
        docs = []
        for name, data in [('manual.txt', b'Pump pressure limit is 42 kPa.'), ('normal.pdf', digital), ('mixed.pdf', mixed), ('table.docx', stream.getvalue())]:
            with patch('services.parser._ocr_page', return_value='Scanned pump pressure limit is 42 kPa.'):
                response = client.post('/api/upload', headers=headers, files={'file': (name, data)})
            assert response.status_code == 200, response.text
            docs.append(response.json())
        assert all(doc['status'] == 'ingested' for doc in client.get('/api/documents', headers=headers).json())
        model = AnswerClient({'status': 'source_excerpts', 'claims': [{'source_id': 'SRC_1', 'quote': 'Pump pressure limit is 42 kPa.'}]})
        with patch.object(query, 'create_llm_client', return_value=model):
            answer = client.post('/api/query', headers=headers, json={'question': 'Pump pressure limit?', 'scope_mode': 'selected', 'document_ids': [docs[0]['id']]}).json()
        assert answer['sources'][0]['document_id'] == docs[0]['id']
        assert client.get('/api/documents/' + docs[0]['id'] + '/file', headers=headers).content == b'Pump pressure limit is 42 kPa.'
        with patch.object(query, 'create_llm_client', return_value=AnswerClient({'status': 'unsupported', 'claims': []})):
            answer = client.post('/api/query', headers=headers, json={'question': 'Compare pump fatalities and dates across documents', 'scope_mode': 'all'}).json()
        assert set(answer['effective_document_ids']) == {d['id'] for d in docs}
        assert answer['evidence_status'] == 'unsupported'
        history = client.get('/api/chat/history', headers=headers).json()
        assert [m['role'] for m in history] == ['user', 'assistant', 'user', 'assistant']
        deleted = client.delete('/api/forget/' + docs[0]['dataset_name'], headers=headers)
        assert deleted.json()['status'] == 'deleted'
        assert client.get('/api/documents/' + docs[0]['id'] + '/file', headers=headers).status_code == 404
        assert client.get('/api/chat/history', headers=headers).json()[1]['sources'][0]['available'] is False
        assert client.post('/api/query', headers=headers, json={'question': 'Pump pressure', 'scope_mode': 'selected', 'document_ids': [docs[0]['id']]}).status_code == 400
